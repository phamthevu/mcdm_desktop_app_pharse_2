import json,re
from datetime import datetime
from PyQt6.QtCore import QDate,QTime,QDateTime,Qt
from PyQt6.QtWidgets import *
from openpyxl import Workbook

LAYOUT={"TITLE","SECTION","DESCRIPTION"}
class TabQuantSurvey(QWidget):
    def __init__(self): super().__init__(); self.db=None; self.config={}; self.widgets={}; self.logged_in_respondent_id=None; self.build_shell()
    def build_shell(self):
        self.root=QVBoxLayout(self); self.access_note=QLabel(); self.access_note.setWordWrap(True); self.root.addWidget(self.access_note)
        top=QHBoxLayout(); self.respondent_label=QLabel("Respondent:"); top.addWidget(self.respondent_label); self.respondent=QComboBox(); top.addWidget(self.respondent,1); self.login_btn=QPushButton("Login..."); self.login_btn.clicked.connect(self.login); top.addWidget(self.login_btn); self.logout_btn=QPushButton("Logout"); self.logout_btn.clicked.connect(self.logout); top.addWidget(self.logout_btn); self.root.addLayout(top)
        self.scroll=QScrollArea(); self.scroll.setWidgetResizable(True); self.formhost=QWidget(); self.form=QVBoxLayout(self.formhost); self.scroll.setWidget(self.formhost); self.root.addWidget(self.scroll,1)
        row=QHBoxLayout(); self.save=QPushButton("SAVE RESPONSE"); self.save.clicked.connect(self.save_response); ex=QPushButton("Export Survey Data..."); ex.clicked.connect(self.export); row.addWidget(self.save); row.addWidget(ex); self.root.addLayout(row)
    def load_project(self, config, db):
        self.config=config; self.db=db; self.logged_in_respondent_id=None; self.rebuild(); self.refresh_respondents_from_db(); self.apply_access_mode()
    def apply_access_mode(self):
        mode=self.config.get('auth_mode','Account')
        free=mode=='Free'
        self.respondent.setVisible(not free); self.respondent_label.setVisible(not free); self.login_btn.setVisible(not free); self.logout_btn.setVisible(not free)
        self.access_note.setText("Access: FREE / ANONYMOUS — no login required; each Save creates a new anonymous submission." if free else "Access: ACCOUNT REQUIRED — log in with the respondent username and password before answering.")
        self.save.setEnabled(free or self.logged_in_respondent_id is not None)
    def login(self):
        if not self.db:return
        dlg=QDialog(self); dlg.setWindowTitle('Respondent Login'); f=QFormLayout(dlg); u=QLineEdit(); pw=QLineEdit(); pw.setEchoMode(QLineEdit.EchoMode.Password); f.addRow('Username',u); f.addRow('Password',pw); bb=QDialogButtonBox(QDialogButtonBox.StandardButton.Ok|QDialogButtonBox.StandardButton.Cancel); bb.accepted.connect(dlg.accept); bb.rejected.connect(dlg.reject); f.addRow(bb)
        if dlg.exec()!=QDialog.DialogCode.Accepted:return
        row=self.db.authenticate_respondent(u.text(),pw.text())
        if not row: QMessageBox.warning(self,'Login Failed','Invalid username or password.'); return
        self.logged_in_respondent_id=row['respondent_id']; i=self.respondent.findData(self.logged_in_respondent_id); self.respondent.setCurrentIndex(i if i>=0 else 0); self.respondent.setEnabled(False); self.login_btn.setEnabled(False); self.logout_btn.setEnabled(True); self.save.setEnabled(True); self.access_note.setText(f"Logged in as: {row.get('full_name') or row.get('username')}")
    def logout(self):
        self.logged_in_respondent_id=None; self.respondent.setEnabled(True); self.login_btn.setEnabled(True); self.logout_btn.setEnabled(False); self.save.setEnabled(False); self.access_note.setText('Access: ACCOUNT REQUIRED — please log in.')
    def reload_from_project_db(self): self.refresh_respondents_from_db()
    def refresh_respondents_from_db(self):
        self.respondent.clear()
        if not self.db:return
        for r in self.db.get_respondents(): self.respondent.addItem(f"{r.get('full_name') or r.get('username')} [{r.get('status')}]",r.get('respondent_id'))
    def rebuild(self):
        while self.form.count():
            x=self.form.takeAt(0); w=x.widget()
            if w:w.deleteLater()
        self.widgets={}
        for q in self.config.get('questions',[]):
            typ=q.get('question_type','').upper(); text=q.get('question_text',''); qid=q.get('question_id')
            if typ in LAYOUT:
                l=QLabel(text); l.setWordWrap(True)
                if typ in {'TITLE','SECTION'}: l.setStyleSheet("font-size:16px;font-weight:bold;margin-top:10px")
                self.form.addWidget(l); continue
            box=QGroupBox(("* " if q.get('required') else "")+text); v=QVBoxLayout(box)
            if q.get('description'): d=QLabel(q['description']); d.setWordWrap(True); d.setStyleSheet("color:#666"); v.addWidget(d)
            w=self.make_widget(q); v.addWidget(w); self.form.addWidget(box); self.widgets[qid]=(w,q,box)
        self.form.addStretch()
    def opts(self,q): return [x.strip() for x in re.split(r'[|;]',q.get('options','')) if x.strip()]
    def make_widget(self,q):
        t=q.get('question_type','').upper(); opts=self.opts(q)
        if t in {'MATRIX_LIKERT','MATRIX_SINGLE_CHOICE'}:
            host=QWidget(); grid=QGridLayout(host); grid.setContentsMargins(0,0,0,0)
            grid.addWidget(QLabel('Statement'),0,0)
            for c,opt in enumerate(opts,1):
                h=QLabel(opt); h.setAlignment(Qt.AlignmentFlag.AlignCenter); grid.addWidget(h,0,c)
            groups={}
            for r,item in enumerate(q.get('matrix_items',[]),1):
                lab=QLabel(item.get('item_text','')); lab.setWordWrap(True); grid.addWidget(lab,r,0)
                group=QButtonGroup(host); group.setExclusive(True)
                for c,opt in enumerate(opts,1):
                    radio=QRadioButton(); radio.setToolTip(opt); group.addButton(radio,c-1); grid.addWidget(radio,r,c,alignment=Qt.AlignmentFlag.AlignCenter)
                groups[item.get('item_id')]=(group,item)
            host._matrix_groups=groups
            host._matrix_options=opts
            return host
        if t=='LONG_TEXT': w=QPlainTextEdit(); w.setMaximumHeight(100); return w
        if t in {'INTEGER','DECIMAL','SHORT_TEXT','EMAIL','PHONE'}:
            w=QLineEdit(); w.setPlaceholderText('Enter a number' if t in {'INTEGER','DECIMAL'} else 'Enter response'); return w
        if t in {'SINGLE_CHOICE','LIKERT','RATING','DROPDOWN','YES_NO','BOOLEAN'}:
            w=QComboBox(); w.addItem("— Select —",None)
            if t in {'YES_NO','BOOLEAN'} and not opts: opts=['Yes','No']
            if t in {'LIKERT','RATING'} and not opts:
                mn=int(q.get('min_value') or 1); mx=int(q.get('max_value') or 5); opts=[str(x) for x in range(mn,mx+1)]
            for x in opts:w.addItem(x,x)
            return w
        if t=='MULTIPLE_CHOICE':
            w=QListWidget(); w.setSelectionMode(QAbstractItemView.SelectionMode.MultiSelection); [w.addItem(x) for x in opts]; w.setMaximumHeight(110); return w
        if t=='DATE': return QDateEdit(QDate.currentDate())
        if t=='TIME': return QTimeEdit(QTime.currentTime())
        if t=='DATETIME': return QDateTimeEdit(QDateTime.currentDateTime())
        return QLineEdit()
    def value(self,w,q):
        t=q.get('question_type','').upper()
        if isinstance(w,QLineEdit): return w.text().strip()
        if isinstance(w,QPlainTextEdit): return w.toPlainText().strip()
        if isinstance(w,QComboBox): return w.currentData()
        if isinstance(w,QListWidget): return [x.text() for x in w.selectedItems()]
        if isinstance(w,QDateEdit): return w.date().toString('yyyy-MM-dd')
        if isinstance(w,QTimeEdit): return w.time().toString('HH:mm:ss')
        if isinstance(w,QDateTimeEdit): return w.dateTime().toString('yyyy-MM-ddTHH:mm:ss')
        if hasattr(w,'_matrix_groups'):
            result={}
            for item_id,(group,_) in w._matrix_groups.items():
                bid=group.checkedId(); result[item_id]=w._matrix_options[bid] if bid >= 0 else None
            return result
    def validate(self):
        errors=[]; data={}
        for qid,(w,q,box) in self.widgets.items():
            box.setStyleSheet(''); val=self.value(w,q); t=q.get('question_type','').upper(); empty=val is None or val=='' or val==[]
            err=''
            if t.startswith('MATRIX_'):
                missing=[item.get('item_text', item_id) for item_id,(group,item) in w._matrix_groups.items() if group.checkedId() < 0]
                if q.get('required') and missing: err='Please answer every row. Missing: ' + '; '.join(missing[:3]) + ('...' if len(missing)>3 else '')
            elif q.get('required') and empty: err='This question is required.'
            if not empty and t in {'INTEGER','DECIMAL'}:
                try:
                    num=int(val) if t=='INTEGER' else float(val)
                    mn=q.get('min_value'); mx=q.get('max_value')
                    if mn is not None and num<float(mn): err=f'Value must be at least {mn}.'
                    if mx is not None and num>float(mx): err=f'Value must be at most {mx}.'
                    if not err: val=num
                except ValueError: err='Value must be an integer.' if t=='INTEGER' else 'Value must be a number.'
            if not empty and t=='EMAIL' and not re.match(r'^[^@\s]+@[^@\s]+\.[^@\s]+$',str(val)): err='Invalid email address.'
            if err: box.setStyleSheet('QGroupBox{border:2px solid #dc3545;}'); errors.append(f"{q.get('question_text')}: {err}")
            data[qid]=val
        return data,errors
    def save_response(self):
        if not self.db:return
        mode=self.config.get('auth_mode','Account')
        if mode=='Account' and not self.logged_in_respondent_id: QMessageBox.warning(self,'Login Required','Please log in before saving a response.'); return
        respondent_id=self.db.create_anonymous_respondent() if mode=='Free' else self.logged_in_respondent_id
        data,errors=self.validate()
        if errors: QMessageBox.warning(self,'Validation Error','Please correct the highlighted fields:\n\n'+'\n'.join(errors[:10])); return
        self.db.save_response(respondent_id,data); self.refresh_respondents_from_db(); QMessageBox.information(self,'Response Saved','Survey response saved successfully.')
    def open_response_for_edit(self,respondent_id):
        i=self.respondent.findData(respondent_id)
        if i>=0:self.respondent.setCurrentIndex(i)
        row=self.db.get_response_by_respondent(respondent_id) if self.db else None
        if not row:return
        data=json.loads(row['response_json'])
        for qid,val in data.items():
            if qid not in self.widgets:continue
            w,q,_=self.widgets[qid]
            if isinstance(w,QLineEdit):w.setText(str(val if val is not None else ''))
            elif isinstance(w,QPlainTextEdit):w.setPlainText(str(val or ''))
            elif isinstance(w,QComboBox):
                j=w.findData(str(val)); w.setCurrentIndex(j if j>=0 else 0)
            elif isinstance(w,QListWidget):
                vals=set(val or [])
                for k in range(w.count()):w.item(k).setSelected(w.item(k).text() in vals)
            elif hasattr(w,'_matrix_groups') and isinstance(val,dict):
                for item_id,item_val in val.items():
                    if item_id not in w._matrix_groups: continue
                    group,_=w._matrix_groups[item_id]
                    try: idx=w._matrix_options.index(str(item_val))
                    except ValueError: continue
                    btn=group.button(idx)
                    if btn: btn.setChecked(True)
    def _export_value(self,data,var_id,questions):
        if var_id in data:
            v=data.get(var_id); return json.dumps(v,ensure_ascii=False) if isinstance(v,list) else v
        for q in questions:
            if str(q.get('question_type','')).startswith('MATRIX_'):
                nested=data.get(q.get('question_id')) or {}
                if var_id in nested: return nested.get(var_id)
        return None
    def export(self):
        if not self.db:return
        p,_=QFileDialog.getSaveFileName(self,'Export Survey Data','survey_results.xlsx','Excel (*.xlsx)')
        if not p:return
        if not p.lower().endswith('.xlsx'):p += '.xlsx'
        wb=Workbook(); meta=wb.active; meta.title='Metadata'; m=self.db.get_meta(); [meta.append([k,v]) for k,v in m.items() if k!='config_json']
        qs=self.config.get('questions',[]); ids=[]
        for q in qs:
            if q.get('question_type') in LAYOUT: continue
            if str(q.get('question_type','')).startswith('MATRIX_'): ids.extend([x.get('item_id') for x in q.get('matrix_items',[])])
            else: ids.append(q.get('question_id'))
        dataws=wb.create_sheet('Data'); dataws.append(['Answer_ID','Respondent_ID','Expert_Name','Username','Source_Project','Submitted_At']+ids)
        for r in self.db.get_all_responses():
            d=json.loads(r['response_json']); dataws.append([f"Ans_{int(r['display_index']):02d}",r['respondent_id'],r.get('full_name'),r.get('username'),r.get('source_project_name'),r.get('submitted_at')]+[json.dumps(d.get(i),ensure_ascii=False) if isinstance(d.get(i),list) else d.get(i) for i in ids])
        cw=wb.create_sheet('Codebook'); cw.append(['Variable_ID','Parent_ID','Question_Text','Question_Type','Required','Options','Min','Max','Description'])
        for q in qs:
            if str(q.get('question_type','')).startswith('MATRIX_'):
                for item in q.get('matrix_items',[]): cw.append([item.get('item_id'),q.get('question_id'),item.get('item_text'),q.get('question_type'),q.get('required'),q.get('options'),q.get('min_value'),q.get('max_value'),q.get('description')])
            else: cw.append([q.get('question_id'),' ',q.get('question_text'),q.get('question_type'),q.get('required'),q.get('options'),q.get('min_value'),q.get('max_value'),q.get('description')])
        rw=wb.create_sheet('Respondent'); rw.append(['Respondent_ID','Username','Full_Name','Organization','Status']); [rw.append([r.get('respondent_id'),r.get('username'),r.get('full_name'),r.get('organization'),r.get('status')]) for r in self.db.get_respondents()]
        wb.save(p); QMessageBox.information(self,'Export Complete',f'Survey data exported successfully.\n\n{p}')
