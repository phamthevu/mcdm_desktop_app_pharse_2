from pathlib import Path
import json
from openpyxl import Workbook
from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import *

class TabQualInterview(QWidget):
    def __init__(self):
        super().__init__(); self.config={}; self.db=None; self.logged_in_respondent_id=None; self.editing_respondent_id=None; self.initial_widgets=[]; self.followups={}
        root=QVBoxLayout(self); root.setContentsMargins(18,14,18,14); root.setSpacing(12)
        self.access=QGroupBox("Participant / Interview Session"); ar=QHBoxLayout(self.access)
        self.free_name=QLineEdit(); self.free_name.setPlaceholderText("Participant name / interview code (required)")
        self.user=QLineEdit(); self.user.setPlaceholderText("Username"); self.password=QLineEdit(); self.password.setEchoMode(QLineEdit.EchoMode.Password); self.password.setPlaceholderText("Password")
        login=QPushButton("Login"); login.clicked.connect(self.login); self.login_state=QLabel("New interview")
        ar.addWidget(QLabel("Participant:")); ar.addWidget(self.free_name,2); ar.addWidget(self.user); ar.addWidget(self.password); ar.addWidget(login); ar.addWidget(self.login_state); root.addWidget(self.access)
        self.scroll=QScrollArea(); self.scroll.setWidgetResizable(True); self.host=QWidget(); self.form=QVBoxLayout(self.host); self.form.setSpacing(12); self.scroll.setWidget(self.host); root.addWidget(self.scroll,1)
        bar=QHBoxLayout(); self.del_btn=QPushButton("Delete selected follow-ups"); self.del_btn.setProperty("danger",True); self.del_btn.clicked.connect(self.delete_selected_followups); self.del_btn.setProperty("secondary",True)
        save=QPushButton("Save Interview & Start New"); save.setProperty("success",True); save.clicked.connect(self.save_response); exp=QPushButton("Export Interviews..."); exp.clicked.connect(self.export)
        bar.addWidget(self.del_btn); bar.addStretch(); bar.addWidget(exp); bar.addWidget(save); root.addLayout(bar)
    def load_project(self,config,db): self.config=config; self.db=db; self.logged_in_respondent_id=None; self.editing_respondent_id=None; self.build_form(); self._apply_access()
    def refresh_respondents_from_db(self): pass
    def _apply_access(self):
        account=str(self.config.get('auth_mode','Free')).strip().lower()=='account'; self.free_name.setVisible(not account); self.user.setVisible(account); self.password.setVisible(account)
    def build_form(self):
        while self.form.count():
            it=self.form.takeAt(0); w=it.widget()
            if w:w.deleteLater()
        self.initial_widgets=[]; self.followups={}
        t=QLabel("In-depth Interview Transcript"); t.setObjectName("pageTitle"); self.form.addWidget(t)
        note=QLabel("Record the initial interview guide below. Use ‘Add follow-up’ under the question being probed; follow-ups remain directly beneath that question as 1.1, 1.2, 2.1 …"); note.setWordWrap(True); note.setObjectName("muted"); self.form.addWidget(note)
        for i,q in enumerate(self.config.get('initial_questions',[]),1):
            box=QGroupBox(f"{i}. {q.get('question_text','')}"); outer=QVBoxLayout(box); a=QPlainTextEdit(); a.setPlaceholderText("Record the participant's answer..."); a.setMinimumHeight(105); outer.addWidget(a)
            add=QPushButton(f"+ Add follow-up after Question {i}"); add.setProperty("secondary",True); outer.addWidget(add,0,Qt.AlignmentFlag.AlignLeft)
            holder=QWidget(); hl=QVBoxLayout(holder); hl.setContentsMargins(22,4,0,0); hl.setSpacing(8); outer.addWidget(holder)
            qid=q.get('question_id'); self.followups[qid]=[]; add.clicked.connect(lambda _=False, qid=qid: self.add_followup(qid)); self.form.addWidget(box); self.initial_widgets.append((q,a,box,hl,i))
        self.form.addStretch()
    def login(self):
        r=self.db.authenticate_respondent(self.user.text(),self.password.text()) if self.db else None
        if not r: QMessageBox.warning(self,"Login Failed","Invalid username or password."); return
        self.logged_in_respondent_id=r['respondent_id']; self.login_state.setText(f"Logged in: {r.get('full_name') or r.get('username')}"); QMessageBox.information(self,"Login Successful","Participant account authenticated successfully.")
    def add_followup(self,parent_id,question="",answer=""):
        row=next((x for x in self.initial_widgets if x[0].get('question_id')==parent_id),None)
        if not row:return
        _,_,_,layout,parent_no=row; idx=len(self.followups[parent_id])+1
        box=QGroupBox(f"{parent_no}.{idx} Follow-up / Probing Question"); v=QVBoxLayout(box); head=QHBoxLayout(); tick=QCheckBox("Select for deletion"); head.addWidget(tick); head.addStretch(); v.addLayout(head)
        q=QLineEdit(); q.setPlaceholderText("Follow-up question"); q.setText(str(question or '')); a=QPlainTextEdit(); a.setPlaceholderText("Participant's answer"); a.setPlainText(str(answer or '')); a.setMinimumHeight(85)
        v.addWidget(QLabel("Question")); v.addWidget(q); v.addWidget(QLabel("Answer")); v.addWidget(a); layout.addWidget(box); self.followups[parent_id].append((q,a,box,tick)); q.setFocus()
    def delete_selected_followups(self):
        selected=[]
        for pid,rows in self.followups.items():
            for item in rows:
                if item[3].isChecked(): selected.append((pid,item))
        if not selected: QMessageBox.information(self,"Delete Follow-ups","Tick one or more follow-up questions first. Initial guide questions cannot be deleted."); return
        if QMessageBox.question(self,"Delete Follow-ups",f"Delete {len(selected)} selected follow-up question(s)?")!=QMessageBox.StandardButton.Yes:return
        for pid,item in selected:
            self.followups[pid].remove(item); item[2].deleteLater()
        self._renumber_followups()
    def _renumber_followups(self):
        for _,_,_,_,parent_no in self.initial_widgets:
            pid=self.initial_widgets[parent_no-1][0].get('question_id')
            for j,(_,_,box,_) in enumerate(self.followups.get(pid,[]),1): box.setTitle(f"{parent_no}.{j} Follow-up / Probing Question")
    def _collect(self):
        items=[]; errors=[]
        for q,w,box,_,parent_no in self.initial_widgets:
            box.setStyleSheet(''); ans=w.toPlainText().strip(); qid=q.get('question_id')
            if q.get('required') and not ans: box.setStyleSheet('QGroupBox{border:2px solid #dc3545;}'); errors.append(q.get('question_text'))
            items.append({'sequence':str(parent_no),'kind':'INITIAL','question_id':qid,'question':q.get('question_text'),'answer':ans})
            for j,(fq,fa,fbox,_) in enumerate(self.followups.get(qid,[]),1):
                fbox.setStyleSheet(''); qt=fq.text().strip(); at=fa.toPlainText().strip()
                if not qt and not at: continue
                if not qt or not at: fbox.setStyleSheet('QGroupBox{border:2px solid #dc3545;}'); errors.append(f"Follow-up {parent_no}.{j}: both question and answer are required"); continue
                items.append({'sequence':f'{parent_no}.{j}','kind':'FOLLOW_UP','parent_question_id':qid,'question_id':None,'question':qt,'answer':at})
        return {'interview_items':items},errors
    def _reset_for_next(self):
        self.editing_respondent_id=None; self.logged_in_respondent_id=None; self.free_name.clear(); self.user.clear(); self.password.clear(); self.login_state.setText("New interview")
        self.build_form(); self._apply_access(); self.scroll.verticalScrollBar().setValue(0); self.free_name.setFocus()
    def save_response(self):
        if not self.db:return
        mode='Account' if str(self.config.get('auth_mode','Free')).strip().lower()=='account' else 'Free'
        if mode=='Account' and not (self.logged_in_respondent_id or self.editing_respondent_id): QMessageBox.warning(self,"Login Required","Please log in before saving a new interview."); return
        name=self.free_name.text().strip() if mode=='Free' else ''
        if mode=='Free' and not name: QMessageBox.warning(self,"Participant Required","Enter the participant name or interview code before recording this interview."); self.free_name.setFocus(); return
        data,errors=self._collect()
        if errors: QMessageBox.warning(self,"Validation Error","Please complete the highlighted interview items:\n\n"+'\n'.join(errors[:10])); return
        rid=self.editing_respondent_id or (self.db.create_free_respondent(name) if mode=='Free' else self.logged_in_respondent_id); self.db.save_response(rid,data)
        QMessageBox.information(self,"Interview Saved",f"Interview saved successfully for {name or self.login_state.text()}.\n\nA clean form will now be prepared for the next participant."); self._reset_for_next()
    def open_response_for_edit(self,respondent_id):
        row=self.db.get_response_by_respondent(respondent_id) if self.db else None
        if not row:return
        self.editing_respondent_id=respondent_id; self.free_name.setText(row.get('full_name') or ''); self.logged_in_respondent_id = respondent_id if str(self.config.get('auth_mode','Account')).strip().lower()=='account' else None; self.login_state.setText(f"Editing existing interview — {row.get('full_name') or row.get('username') or respondent_id}")
        items=json.loads(row['response_json']).get('interview_items',[]); initial={x.get('question_id'):x for x in items if x.get('kind')=='INITIAL'}
        for q,w,_,_,_ in self.initial_widgets:w.setPlainText(str(initial.get(q.get('question_id'),{}).get('answer','')))
        for x in items:
            if x.get('kind')=='FOLLOW_UP':
                pid=x.get('parent_question_id')
                if not pid:
                    # Legacy V4 interviews had all follow-ups at the end; attach to the last initial question when editing.
                    pid=self.initial_widgets[-1][0].get('question_id') if self.initial_widgets else None
                if pid:self.add_followup(pid,str(x.get('question','')),str(x.get('answer','')))
    def export(self):
        if not self.db:return
        default_name=f"{self.db.get_meta().get('project_name','interview_results')}_Interview.xlsx"
        name,ok=QInputDialog.getText(self,"Export Qualitative Interviews","Output file name:",text=default_name)
        if not ok or not name.strip():return
        name=name.strip()
        if not name.lower().endswith('.xlsx'):name+='.xlsx'
        folder=QFileDialog.getExistingDirectory(self,"Choose Export Folder")
        if not folder:return
        p=str(Path(folder)/name)
        wb=Workbook(); meta=wb.active; meta.title='Metadata'; m=self.db.get_meta(); [meta.append([k,v]) for k,v in m.items() if k!='config_json']; lookup=wb.create_sheet('Lookup'); lookup.append(['Interview_ID','Respondent_ID','Participant_Name','Username','Source_Project','Submitted_At','Total_Items'])
        for r in self.db.get_all_responses():
            iid=f"Interview_{int(r['display_index']):02d}"; items=json.loads(r['response_json']).get('interview_items',[]); lookup.append([iid,r.get('respondent_id'),r.get('full_name'),r.get('username'),r.get('source_project_name'),r.get('submitted_at'),len(items)]); ws=wb.create_sheet(iid[:31]); ws.append(['Sequence','Type','Question_ID','Question','Answer'])
            for x in items:ws.append([x.get('sequence'),x.get('kind'),x.get('question_id'),x.get('question'),x.get('answer')])
            ws.column_dimensions['D'].width=60; ws.column_dimensions['E'].width=80
        guide=wb.create_sheet('Initial_Guide'); guide.append(['Question_ID','Question_Text','Required','Display_Order'])
        for q in self.config.get('initial_questions',[]):guide.append([q.get('question_id'),q.get('question_text'),q.get('required'),q.get('display_order')])
        wb.save(p); QMessageBox.information(self,"Export Complete",f"Qualitative interviews exported successfully.\n\n{p}")
