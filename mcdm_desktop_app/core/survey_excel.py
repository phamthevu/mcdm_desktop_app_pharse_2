from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from core.survey_models import SurveyQuestion, SUPPORTED_TYPES

COLS = ["Question_ID","Question_Text","Question_Type","Required","Options","Min_Value","Max_Value","Default_Value","Validation","Description","Display_Order"]
MATRIX_COLS = ["Parent_ID", "Item_ID", "Item_Text", "Display_Order"]

def _bool(v): return str(v).strip().lower() in {"1","true","yes","y","required"}
def _clean(v): return "" if pd.isna(v) else str(v).strip()
def _num(v):
    if pd.isna(v) or str(v).strip()=="": return None
    return float(v)

def parse_survey_excel(path):
    xl = pd.ExcelFile(path)
    qsheet = "Questions" if "Questions" in xl.sheet_names else xl.sheet_names[0]
    df = pd.read_excel(path, sheet_name=qsheet)
    missing = [c for c in ["Question_ID","Question_Text","Question_Type"] if c not in df.columns]
    if missing: raise ValueError("Missing required columns: " + ", ".join(missing))

    matrix_map = {}
    if "Matrix_Items" in xl.sheet_names:
        mdf = pd.read_excel(path, sheet_name="Matrix_Items")
        mmissing = [c for c in ["Parent_ID", "Item_ID", "Item_Text"] if c not in mdf.columns]
        if mmissing: raise ValueError("Matrix_Items is missing: " + ", ".join(mmissing))
        seen_items = set()
        for i, r in mdf.iterrows():
            parent, item_id, item_text = _clean(r.get("Parent_ID")), _clean(r.get("Item_ID")), _clean(r.get("Item_Text"))
            if not parent and not item_id and not item_text: continue
            if not parent or not item_id or not item_text: raise ValueError(f"Incomplete Matrix_Items row {i+2}.")
            if item_id in seen_items: raise ValueError(f"Duplicate matrix Item_ID: {item_id}")
            seen_items.add(item_id)
            order = int(_num(r.get("Display_Order")) or i+1)
            matrix_map.setdefault(parent, []).append({"item_id": item_id, "item_text": item_text, "display_order": order})
        for items in matrix_map.values(): items.sort(key=lambda x: x["display_order"])

    questions=[]; ids=set()
    for i,row in df.iterrows():
        qid=_clean(row.get("Question_ID")); qt=_clean(row.get("Question_Type")).upper()
        if not qid and qt not in {"TITLE","SECTION","DESCRIPTION"}: continue
        if not qid: qid=f"LAYOUT_{i+1}"
        if qid in ids: raise ValueError(f"Duplicate Question_ID: {qid}")
        ids.add(qid)
        if qt not in SUPPORTED_TYPES: raise ValueError(f"Unsupported Question_Type '{qt}' at {qid}")
        mn=_num(row.get("Min_Value")); mx=_num(row.get("Max_Value"))
        if mn is not None and mx is not None and mn > mx: raise ValueError(f"Min_Value > Max_Value at {qid}")
        items = matrix_map.get(qid, [])
        if qt.startswith("MATRIX_") and not items: raise ValueError(f"{qid} is {qt} but has no rows in Matrix_Items.")
        questions.append(SurveyQuestion(qid,_clean(row.get("Question_Text")),qt,_bool(row.get("Required")),_clean(row.get("Options")),mn,mx,_clean(row.get("Default_Value")),_clean(row.get("Validation")),_clean(row.get("Description")),int(_num(row.get("Display_Order")) or i+1),items))
    questions.sort(key=lambda q:q.display_order)

    unknown_parents = sorted(set(matrix_map) - {q.question_id for q in questions})
    if unknown_parents: raise ValueError("Matrix_Items references unknown Parent_ID: " + ", ".join(unknown_parents))

    accounts=[]
    for sheet in ("Accounts","Experts","Respondents"):
        if sheet in xl.sheet_names:
            adf=pd.read_excel(path,sheet_name=sheet)
            for _,r in adf.iterrows():
                u=_clean(r.get("Username"))
                if not u: continue
                accounts.append({"Respondent_ID":_clean(r.get("Respondent_ID")),"Username":u,"Password":_clean(r.get("Password")) or "123","Full_Name":_clean(r.get("Full_Name")) or _clean(r.get("Expert_Name")),"Organization":_clean(r.get("Organization"))})
            break
    return questions, accounts

def generate_survey_template(path):
    wb=Workbook(); ws=wb.active; ws.title="Questions"; ws.append(COLS)
    samples=[
      ["INTRO","Research Survey","TITLE",False,"","","","","","Please answer all applicable questions.",1],
      ["NAME","Full name","SHORT_TEXT",True,"","","","","","",2],
      ["AGE","Age","INTEGER",True,"",18,100,"","","Enter your age in years.",3],
      ["GENDER","Gender","SINGLE_CHOICE",True,"Male|Female|Other","","","","","",4],
      ["SAT","Satisfaction with ChatGPT","MATRIX_LIKERT",True,"1|2|3|4|5",1,5,"","","1 = Strongly disagree; 5 = Strongly agree",5],
      ["INCOME","Monthly income","DECIMAL",False,"",0,1000000000,"","","Numeric value only",6],
      ["COMMENT","Additional comments","LONG_TEXT",False,"","","","","","",7],
    ]
    for x in samples: ws.append(x)
    mi=wb.create_sheet("Matrix_Items"); mi.append(MATRIX_COLS)
    for x in [
        ["SAT","SAT1","ChatGPT is easy to use",1],
        ["SAT","SAT2","ChatGPT saves me time",2],
        ["SAT","SAT3","ChatGPT improves my work",3],
        ["SAT","SAT4","I am satisfied with ChatGPT",4],
    ]: mi.append(x)
    ac=wb.create_sheet("Accounts"); ac.append(["Respondent_ID","Username","Password","Full_Name","Organization"]); ac.append(["R001","expert01","123","Expert 01","Organization A"])
    for sh in wb.worksheets:
        sh.freeze_panes="A2"; sh.auto_filter.ref=sh.dimensions
        for col in sh.columns:
            sh.column_dimensions[col[0].column_letter].width=min(max(len(str(c.value or "")) for c in col)+2,45)
    Path(path).parent.mkdir(parents=True,exist_ok=True); wb.save(path)
